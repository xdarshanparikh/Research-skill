"""Small python-pptx helpers so every chart slide carries the same source footer and notes.

    from pptx import Presentation
    from pptx_helpers import add_chart_slide, add_source, add_notes, LAYOUT
    prs = Presentation("template.pptx")            # or Presentation(); set 16:9 with widescreen(prs)
    s = add_chart_slide(prs, "Imports more than doubled since 2019",
                        "charts/imports.png", source="Source: IEA, Gas Market Report, 2026",
                        subtitle="LNG imports, million tonnes", notes="Derived: 2025 shares applied to FY26.")
    prs.save("deck.pptx")

Footer rules (check_deck.py enforces them): one text box named "Source" per content slide, same
position on every slide, 1-2 lines, publisher + document + year. Full links go in the speaker
notes and research/Sources.md, not on the slide.
"""
from pptx.util import Emu, Inches, Pt
from pptx.dml.color import RGBColor

# positions in inches for a 13.333 x 7.5 slide; scaled automatically for other sizes
LAYOUT = {"margin": 0.6, "title_top": 0.4, "title_h": 0.9, "sub_h": 0.45, "chart_top": 1.85,
          "src_top": 6.85, "src_h": 0.45, "title_pt": 30, "sub_pt": 16, "src_pt": 10}


def widescreen(prs):
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def _k(prs):
    return prs.slide_width / Inches(13.333)


def _text(slide, name, x, y, w, h, text, pt, color="1A1A1A", bold=False, italic=False):
    tb = slide.shapes.add_textbox(x, y, w, h); tb.name = name
    tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = line
        r.font.size = Pt(pt); r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = RGBColor.from_string(color)
    return tb


def add_source(slide, prs, text, color="666666"):
    """The one standard footer. Replaces an existing "Source" box so re-runs don't stack footers."""
    for sh in list(slide.shapes):
        if sh.name == "Source": sh._element.getparent().remove(sh._element)
    k, L = _k(prs), LAYOUT
    return _text(slide, "Source", Inches(L["margin"] * k), Inches(L["src_top"] * k),
                 prs.slide_width - Inches(2 * L["margin"] * k), Inches(L["src_h"] * k), text, L["src_pt"], color)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text
    return slide


def add_picture_fit(slide, prs, png, top_in, bottom_in, margin_in=None):
    """Place a chart PNG as large as fits between top and bottom, centred, keeping its aspect ratio."""
    from PIL import Image
    k = _k(prs); m = Inches((margin_in if margin_in is not None else LAYOUT["margin"]) * k)
    box_w = prs.slide_width - 2 * m; box_h = Inches((bottom_in - top_in) * k)
    w, h = Image.open(png).size
    scale = min(box_w / w, box_h / h)
    pw, ph = Emu(int(w * scale)), Emu(int(h * scale))
    pic = slide.shapes.add_picture(str(png), m + (box_w - pw) // 2, Inches(top_in * k), pw, ph)
    pic.name = "Chart"
    return pic


def add_chart_slide(prs, title, png, source, subtitle=None, notes=None, layout_index=6):
    """Title (the takeaway, not the topic) + optional subtitle (what is measured, unit, period) +
    chart picture + source footer + speaker notes. layout_index 6 = Blank in the default template."""
    k, L = _k(prs), LAYOUT
    slide = prs.slides.add_slide(prs.slide_layouts[min(layout_index, len(prs.slide_layouts) - 1)])
    full_w = prs.slide_width - Inches(2 * L["margin"] * k)
    _text(slide, "Title", Inches(L["margin"] * k), Inches(L["title_top"] * k), full_w, Inches(L["title_h"] * k), title, L["title_pt"], bold=True)
    if subtitle:
        _text(slide, "Subtitle", Inches(L["margin"] * k), Inches((L["title_top"] + L["title_h"]) * k), full_w, Inches(L["sub_h"] * k), subtitle, L["sub_pt"], "555555", italic=True)
    add_picture_fit(slide, prs, png, L["chart_top"], L["src_top"] - 0.1)
    add_source(slide, prs, source)
    if notes: add_notes(slide, notes)
    return slide
