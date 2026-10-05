"""Pre-flight check for a research deck (.pptx).

usage: python check_deck.py deck.pptx [--skip 1,12] [--max-age 2] [--today 2026-10-05]

Flags, per slide:
- NO SOURCE      a content slide with a number, chart, table or picture but no "Source" footer
- FOOTER MOVED   the footer is not in the same place as on the other slides
- LONG FOOTER    more than two lines (links belong in notes and Sources.md)
- OLD YEAR       the newest year in the footer is older than --max-age years; mark it checked
                 with "(latest available)" in the footer or "data-checked: YYYY" in the notes
- OFF SLIDE      a shape runs past the slide edge
- OVERLAP        a text box overlaps a chart, picture, table or another text box
Skip title/section/closing slides with --skip (1-based). Exit 1 when anything is flagged.
"""
import datetime
import re
import sys

from pptx import Presentation
from pptx.util import Emu

YEAR = re.compile(r"\b(19[5-9]\d|20[0-4]\d)\b")


def box(sh):
    return (sh.left or 0, sh.top or 0, (sh.left or 0) + (sh.width or 0), (sh.top or 0) + (sh.height or 0))


def overlap(a, b, tol=Emu(45720)):  # 0.05 in: headings sit just above their charts by design
    return a[0] + tol < b[2] and b[0] + tol < a[2] and a[1] + tol < b[3] and b[1] + tol < a[3]


def kind(sh):
    if sh.name == "Source": return "source"
    if sh.shape_type == 13 or sh.name.startswith("Chart"): return "visual"
    if getattr(sh, "has_chart", False) and sh.has_chart: return "visual"
    if getattr(sh, "has_table", False) and sh.has_table: return "visual"
    if sh.has_text_frame and sh.text_frame.text.strip(): return "text"
    return "other"


def main():
    args = sys.argv[1:]
    path = args[0]
    opt = lambda k, d: args[args.index(k) + 1] if k in args else d
    skip = {int(x) for x in opt("--skip", "").split(",") if x}
    max_age = int(opt("--max-age", "2"))
    today = datetime.date.fromisoformat(opt("--today", datetime.date.today().isoformat()))
    prs = Presentation(path); W, H = prs.slide_width, prs.slide_height
    issues, footers = [], []
    for i, slide in enumerate(prs.slides, 1):
        shapes = list(slide.shapes)
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        src = [s for s in shapes if s.name == "Source"]
        body = " ".join(s.text_frame.text for s in shapes if s.has_text_frame and s.name != "Source")
        needs = any(kind(s) == "visual" for s in shapes) or bool(re.search(r"\d", body))
        if i not in skip:
            if needs and not src:
                issues.append((i, "NO SOURCE", "add the standard footer (pptx_helpers.add_source)"))
            for s in src:
                footers.append((i, s.left, s.top))
                t = s.text_frame.text
                lines = [l for l in t.split("\n") if l.strip()]
                if len(lines) > 2 or len(t) > 260:
                    issues.append((i, "LONG FOOTER", f"{len(lines)} lines / {len(t)} chars"))
                ys = [int(y) for y in YEAR.findall(t)]
                checked = "latest available" in t.lower() or re.search(r"data-checked:\s*\d{4}", notes)
                if ys and today.year - max(ys) > max_age and not checked:
                    issues.append((i, "OLD YEAR", f"newest year {max(ys)}: look for a newer edition, or mark it checked"))
        for s in shapes:
            l, t, r, b = box(s)
            if l < 0 or t < 0 or r > W + Emu(5000) or b > H + Emu(5000):
                issues.append((i, "OFF SLIDE", s.name))
        tv = [s for s in shapes if kind(s) in ("text", "visual", "source")]
        for a in range(len(tv)):
            for c in range(a + 1, len(tv)):
                if kind(tv[a]) == kind(tv[c]) == "visual": continue
                if overlap(box(tv[a]), box(tv[c])):
                    issues.append((i, "OVERLAP", f"{tv[a].name} / {tv[c].name}"))
    if footers:
        ref = max(set((l, t) for _, l, t in footers), key=lambda p: sum(1 for _, l, t in footers if (l, t) == p))
        for i, l, t in footers:
            if abs(l - ref[0]) > Emu(30000) or abs(t - ref[1]) > Emu(30000):
                issues.append((i, "FOOTER MOVED", "use the standard position"))
    for i, k, m in sorted(issues):
        print(f"slide {i:2d}  {k:13s} {m}")
    print("all clear" if not issues else f"\n{len(issues)} issue(s)")
    sys.exit(1 if issues else 0)


if __name__ == "__main__":
    main()
