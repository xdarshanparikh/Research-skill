"""Build slides in the MGL template look with python-pptx.

    import sys; sys.path.insert(0, "<skill>/scripts")
    from mgl_kit import Deck
    d = Deck()                                            # opens <skill>/template/mgl-template.pptx
    d.title_slide("5-Year Business Plan, FY27-FY32", "External context", "25th Sep 2026")
    s = d.content_slide("Global gas to FY32: a short squeeze, then a long surplus",
                        "What the latest outlooks say: BloombergNEF (Jul 2026), Shell (Jun 2026)")
    d.heading(s, 0.9, 1.42, 7.5, "Global LNG surplus / (shortage), million tonnes")
    d.column_chart(s, 0.81, 1.67, 7.57, 2.15, ["FY27", "FY28"], {"Base": [0, -3], "Delay": [0, -17]})
    d.stat_card(s, 8.71, 1.42, 3.72, 1.0, "~20%", "of global LNG supply still offline", icon="waves", tone="red")
    d.takeaway(s, "FY27 is the peak-pain year. Prices roughly halve in FY28.")
    d.sources(s, "Sources: BloombergNEF, Global LNG Outlook, 2026; Shell, LNG Outlook, 2026")
    d.notes(s, "Full links and derivations here.")
    d.closing_slide()
    d.save("deck.pptx")

All positions are in inches on the 13.33 x 7.5 slide and follow the template's grid (see
references/template-design.md). Content sits between x 0.9 and 12.43, below the subtitle (y ~1.35) and above
the takeaway bar (y ~6.0); the source line is at y 6.68. check_deck.py reads the footer named "Source".
"""
import copy
import pathlib

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SKILL = pathlib.Path(__file__).resolve().parent.parent
TEMPLATE = SKILL / "template" / "mgl-template.pptx"
ICONS = SKILL / "template" / "icons"

# ---- design tokens, taken from the template's own slides -------------------------------------
C = {
    "green": "51AD30",       # brand green: titles, positive icons, highlight bars
    "dark_green": "2F7A17",  # positive numbers and text on light green
    "blue": "0067B2",        # "Key takeaway" label, primary series, emphasis text
    "light_blue": "7FB2D9",
    "amber": "FCB000",       # second series, "watch"
    "amber_text": "9A6A00",
    "red": "B9382A",         # risk, negative, peak-pain numbers
    "ink": "000000",         # headings, body on cards
    "text": "404040",        # subtitles, card labels, axis labels
    "muted": "6E6E6E",       # source line
    "axis": "888888",
    "grid": "EDEDED",
    "card": "F5F5F5",        # neutral card / panel
    "green_tint": "EAF5E6",  # takeaway bar, positive card, table header
    "amber_tint": "FFF3D6",  # "watch" callout
    "red_tint": "F6E1DE",    # "risk" callout / bad cell
    "white": "FFFFFF",
}
# chart series order (validated: CVD separation >= 7 with labels, normal-vision PASS); last = context grey
SERIES = ["0067B2", "FCB000", "51AD30", "B9382A", "7FB2D9", "888888"]
FONT = "Aptos"  # the template's slide font; Office 365 has it. Use "Calibri" for older Office.
TONES = {  # stat-card tone -> (circle fill, number colour)
    "green": (C["green"], C["dark_green"]), "red": (C["red"], C["red"]),
    "blue": (C["blue"], C["blue"]), "amber": (C["amber"], C["amber_text"]), "ink": (C["green"], C["ink"]),
}
CALLOUT = {"watch": (C["amber_tint"], C["amber_text"]), "risk": (C["red_tint"], C["red"]),
           "good": (C["green_tint"], C["dark_green"]), "note": (C["card"], C["blue"])}
CELL = {"good": (C["green_tint"], C["dark_green"]), "bad": (C["red_tint"], C["red"]),
        "watch": (C["amber_tint"], C["amber_text"]), None: (C["card"], C["ink"])}
GRID = {"left": 0.9, "right": 12.43, "width": 11.53, "sub_y": 0.95, "body_y": 1.42,
        "takeaway_y": 6.0, "takeaway_h": 0.6, "src_y": 6.68, "src_h": 0.28}


def rgb(h):
    return RGBColor.from_string(h)


class Deck:
    def __init__(self, template=TEMPLATE):
        self.prs = Presentation(str(template))
        self.L = {l.name: l for l in self.prs.slide_layouts}

    # ---------------------------------------------------------------- slides
    def title_slide(self, title, subtitle="", date=""):
        s = self.prs.slides.add_slide(self.L["Title Slide"])
        t = s.shapes.title; t.text = title
        t.left, t.top, t.width, t.height = 0, Inches(2.37), Inches(9.8), Inches(1.09)
        sub = [p for p in s.placeholders if p.placeholder_format.idx == 1]
        if sub:
            b = sub[0]; b.text = subtitle
            b.left, b.top, b.width, b.height = 0, Inches(3.48), Inches(9.8), Inches(0.55)
            if date:  # the template places the date as a second subtitle box, lower down
                el = copy.deepcopy(b._element); s.shapes._spTree.append(el)
                box = s.shapes[-1]; box.top = Inches(4.5); box.height = Inches(0.73); box.text_frame.text = date
        self._number(s)
        return s

    def section_slide(self, title, text=""):
        s = self.prs.slides.add_slide(self.L["Section Header"])
        s.shapes.title.text = title
        for p in s.placeholders:
            if p.placeholder_format.idx == 1:
                if text: p.text = text
                else: p._element.getparent().remove(p._element)
        self._number(s)
        return s

    def content_slide(self, title, subtitle=None):
        """Green takeaway title at the top (template: 24 pt, brand green) + grey subtitle line."""
        s = self.prs.slides.add_slide(self.L["Title Only"])
        s.shapes.title.text = title
        for r in s.shapes.title.text_frame.paragraphs[0].runs:
            r.font.size = Pt(24)
        if subtitle:
            self.text(s, GRID["left"], GRID["sub_y"], GRID["width"], 0.33, subtitle, 12, C["text"], name="Subtitle")
        self._number(s)
        return s

    def closing_slide(self):
        s = self.prs.slides.add_slide(self.L["Custom Layout"])  # green "Thank You..." panel + address
        return s

    def save(self, path):
        self.prs.save(str(path))
        return path

    def _number(self, s):
        """Copy the layout's slide-number placeholder onto the slide so the green page dot shows a number."""
        for ph in s.slide_layout.placeholders:
            if ph.placeholder_format.type is not None and "SLIDE_NUMBER" in str(ph.placeholder_format.type):
                if not any(p.placeholder_format.idx == ph.placeholder_format.idx for p in s.placeholders):
                    el = copy.deepcopy(ph._element)
                    for t in el.iter(qn("a:t")): t.text = "‹#›"
                    s.shapes._spTree.append(el)

    # ---------------------------------------------------------------- text blocks
    def text(self, s, x, y, w, h, text, size=11, color=C["ink"], bold=False, italic=False, align="l",
             anchor="t", name=None, runs=None):
        """One text box, zero padding so it lines up with cards. `runs` = [(text, {bold, color, size})] for mixed styles;
        a list of strings in `text` makes one paragraph each."""
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name: tb.name = name
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        paras = runs and [runs] or [[(t, {})] for t in (text if isinstance(text, list) else [text])]
        for i, pr in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            for t, o in pr:
                r = p.add_run(); r.text = t; f = r.font
                f.name = FONT; f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold); f.italic = o.get("italic", italic)
                f.color.rgb = rgb(o.get("color", color))
        return tb

    def heading(self, s, x, y, w, text, size=11):
        """Bold black chart / block heading (template: 11 pt bold, just above its chart)."""
        return self.text(s, x, y, w, 0.28, text, size, C["ink"], bold=True)

    def bullets(self, s, x, y, w, h, items, size=10, color=C["ink"]):
        """Bulleted list; an item may be (lead, rest) to bold its lead-in ("City gas leads: ...")."""
        tb = self.text(s, x, y, w, h, "", size, color)
        tf = tb.text_frame
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            pPr = p._p.get_or_add_pPr(); pPr.set("marL", "171450"); pPr.set("indent", "-171450")
            bu = pPr.makeelement(qn("a:buChar"), {"char": "•"}); pPr.append(bu)
            p.space_after = Pt(4)
            lead, rest = it if isinstance(it, tuple) else ("", it)
            for t, b in ((lead, True), (rest, False)):
                if not t: continue
                r = p.add_run(); r.text = (t + " ") if b else t
                r.font.name = FONT; r.font.size = Pt(size); r.font.bold = b; r.font.color.rgb = rgb(color)
        return tb

    def rect(self, s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, name=None):
        r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        r.fill.solid(); r.fill.fore_color.rgb = rgb(fill)
        if line: r.line.color.rgb = rgb(line); r.line.width = Pt(1)
        else: r.line.fill.background()
        r.shadow.inherit = False
        if name: r.name = name
        return r

    def icon(self, s, x, y, d, name, fill=C["green"]):
        """White pictogram in a filled circle (template motif). d = circle diameter in inches."""
        self.rect(s, x, y, d, d, fill, shape=MSO_SHAPE.OVAL, name=f"Icon circle {name}")
        p = ICONS / f"{name}.png"
        if p.exists():
            k = d * 0.52
            s.shapes.add_picture(str(p), Inches(x + (d - k) / 2), Inches(y + (d - k) / 2), Inches(k), Inches(k))

    # ---------------------------------------------------------------- components
    def stat_card(self, s, x, y, w, h, value, label, icon=None, tone="green", fill="card", value_size=None):
        """Grey (or green-tint) card: icon circle left, big coloured number, small grey line.
        Template sizes: 3.72 x 1.0 in, number 18-19 pt bold, label 9 pt."""
        self.rect(s, x, y, w, h, C["green_tint"] if fill == "green" else C["card"], name="Card")
        circ, numc = TONES[tone]
        tx = x + 0.14
        if icon:
            d = min(0.52, h * 0.55); self.icon(s, x + 0.14, y + (h - d) / 2, d, icon, circ); tx = x + 0.14 + d + 0.13
        tw = x + w - tx - 0.1
        vs = value_size or (19 if h >= 0.95 else 17)
        self.text(s, tx, y + 0.07, tw, h * 0.4, value, vs, numc if fill != "green" else C["ink"], bold=True, anchor="m")
        self.text(s, tx, y + 0.07 + h * 0.4, tw, h * 0.5, label, 9 if h >= 0.95 else 8.5, C["text"])

    def stat_row(self, s, items, y=1.33, h=0.84, x=None, w=None, gap=0.17, fill="green"):
        """A row of stat cards across the content width. items = [(value, label, icon, tone)]."""
        x = GRID["left"] if x is None else x; w = GRID["width"] if w is None else w
        cw = (w - gap * (len(items) - 1)) / len(items)
        for i, it in enumerate(items):
            v, l, ic, tone = (list(it) + [None, "green"])[:4]
            self.stat_card(s, x + i * (cw + gap), y, cw, h, v, l, ic, tone or "green", fill=fill, value_size=17)

    def info_card(self, s, x, y, w, h, title, body, highlight=None, fill="card", icon=None, tone="green", tag=None):
        """Panel with a bold title, an optional blue highlight line and grey body (template: GA cards, scenario cards).
        tag = ("Tailwind", "good"|"watch"|"bad"|"note") adds a status pill at the top right."""
        self.rect(s, x, y, w, h, C["green_tint"] if fill == "green" else C["card"], name="Card")
        tx = x + 0.14
        if icon:
            self.icon(s, x + 0.12, y + 0.12, 0.42, icon, TONES[tone][0]); tx = x + 0.66
        tw = x + w - tx - 0.12; yy = y + 0.08
        if tag:
            self.tag(s, x + w - 0.97, y + 0.1, 0.85, tag[0], {"bad": "risk"}.get(tag[1], tag[1]))
        self.text(s, tx, yy, tw - (0.95 if tag else 0), 0.28, title, 10, C["ink"], bold=True); yy += 0.3
        if highlight:
            self.text(s, tx, yy, tw, 0.3, highlight, 9, C["blue"], bold=True); yy += 0.32
        self.text(s, tx, yy, tw, y + h - yy - 0.06, body, 8.5, C["text"])

    def callout(self, s, x, y, w, h, lead, text, kind="watch"):
        """Tinted box with a coloured bold lead-in: watch (amber), risk (red), good (green), note (grey/blue)."""
        fill, col = CALLOUT[kind]
        self.rect(s, x, y, w, h, fill, name="Callout")
        self.text(s, x + 0.12, y + 0.06, w - 0.24, h - 0.12, "", 9,
                  runs=[(lead + " ", {"bold": True, "color": col}), (text, {"color": C["ink"]})])

    def tag(self, s, x, y, w, text, kind="good"):
        """Small outlined status pill ("High risk", "Tailwind", "Positive")."""
        fill, col = CALLOUT[kind]
        r = self.rect(s, x, y, w, 0.22, fill, line=col, name="Tag")
        self.text(s, x, y, w, 0.22, text, 7.5, col, bold=True, align="c", anchor="m")
        return r

    def takeaway(self, s, text, y=None, h=None, label="Key takeaway"):
        """Light-green bar near the bottom: blue bold 'Key takeaway' + one or two sentences (template: 11 pt)."""
        y = GRID["takeaway_y"] if y is None else y; h = GRID["takeaway_h"] if h is None else h
        self.rect(s, GRID["left"], y, GRID["width"], h, C["green_tint"], name="Takeaway")
        self.text(s, GRID["left"] + 0.2, y + 0.05, GRID["width"] - 0.4, h - 0.1, "", 11, anchor="m", name="Takeaway text",
                  runs=[(label + "  ", {"bold": True, "color": C["blue"]}), (text, {"color": C["ink"]})])

    def sources(self, s, text, y=None):
        """The one standard footer, named "Source" (check_deck.py looks for it): 7.5-8 pt grey, under the takeaway."""
        for sh in list(s.shapes):
            if sh.name == "Source": sh._element.getparent().remove(sh._element)
        return self.text(s, GRID["left"], GRID["src_y"] if y is None else y, 11.03, GRID["src_h"], text, 7.5, C["muted"], name="Source")

    def notes(self, s, text):
        s.notes_slide.notes_text_frame.text = text

    def picture(self, s, png, x, y, w=None, h=None):
        """Place a charts.py PNG (sankey, waterfall, slope...) keeping its aspect ratio inside w x h."""
        from PIL import Image
        iw, ih = Image.open(png).size
        if w and h:
            k = min(w / iw, h / ih); pw, ph = iw * k, ih * k
            x, y = x + (w - pw) / 2, y + (h - ph) / 2
        elif w: pw, ph = w, w * ih / iw
        else: pw, ph = h * iw / ih, h
        pic = s.shapes.add_picture(str(png), Inches(x), Inches(y), Inches(pw), Inches(ph)); pic.name = "Chart image"
        return pic

    # ---------------------------------------------------------------- tables
    def table(self, s, x, y, w, rows, col_w=None, row_h=0.36, size=8.5, header=True):
        """rows = list of lists; a cell may be "text" or (text, "good"|"bad"|"watch") to tint it like the
        template's scorecards. Header row: light green, bold. Body: light grey with white 2.5 pt gaps."""
        nr, nc = len(rows), max(len(r) for r in rows)
        gf = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
        gf.name = "Table"; t = gf.table
        tblPr = t._tbl.tblPr
        for a in ("bandRow", "firstRow"): tblPr.set(a, "0")
        sid = tblPr.find(qn("a:tableStyleId"))
        if sid is not None: sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # "No Style, No Grid"
        if col_w:
            for i, cw in enumerate(col_w): t.columns[i].width = Inches(cw)
        for ri, r in enumerate(rows):
            t.rows[ri].height = Inches(row_h)
            for ci in range(nc):
                v = r[ci] if ci < len(r) else ""
                txt, kind = (v if isinstance(v, tuple) else (v, None))
                cell = t.cell(ri, ci)
                fill, col = (C["green_tint"], C["ink"]) if (header and ri == 0) else CELL[kind]
                cell.fill.solid(); cell.fill.fore_color.rgb = rgb(fill)
                cell.margin_left = cell.margin_right = Inches(0.06); cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame; tf.word_wrap = True
                p = tf.paragraphs[0]; p.text = ""
                run = p.add_run(); run.text = str(txt)
                run.font.name = FONT; run.font.size = Pt(size); run.font.color.rgb = rgb(col)
                run.font.bold = (header and ri == 0) or ci == 0 or kind is not None
                p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
                _cell_borders(cell)
        return gf

    # ---------------------------------------------------------------- native charts
    def column_chart(self, s, x, y, w, h, cats, series, stacked=False, horizontal=False, colors=None,
                     point_colors=None, fmt='#,##0', labels=True, legend=None, line=None, gap=45):
        """Clustered / stacked column or bar chart in the template style: data labels on, no gridlines,
        grey axis line, Aptos 9-10 pt, legend at the bottom when there are several series.
        point_colors = list of hex per category for a single-series highlight (e.g. FY27 red, rest green).
        line = {"name": [values]} overlays a line series (e.g. demand vs contracted supply)."""
        cd = CategoryChartData(); cd.categories = cats
        for n, v in series.items(): cd.add_series(n, v)
        kind = {(False, False): XL_CHART_TYPE.COLUMN_CLUSTERED, (True, False): XL_CHART_TYPE.COLUMN_STACKED,
                (False, True): XL_CHART_TYPE.BAR_CLUSTERED, (True, True): XL_CHART_TYPE.BAR_STACKED}[(stacked, horizontal)]
        gf = s.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd); gf.name = "Chart"
        ch = gf.chart; ch.has_title = False
        cols = colors or SERIES
        plot = ch.plots[0]; plot.gap_width = gap
        if stacked: plot.overlap = 100
        for i, ser in enumerate(plot.series):
            ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(cols[i % len(cols)])
            ser.invert_if_negative = False
        if point_colors and len(plot.series) == 1:
            for i, pc in enumerate(point_colors):
                pt = plot.series[0].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(pc)
        if labels:
            plot.has_data_labels = True; dl = plot.data_labels
            dl.number_format = fmt; dl.number_format_is_linked = False
            dl.font.size = Pt(9 if stacked else 10); dl.font.name = FONT
            dl.font.color.rgb = rgb(C["white"] if stacked else C["ink"])
            dl.position = XL_LABEL_POSITION.CENTER if stacked else XL_LABEL_POSITION.OUTSIDE_END
        _style_axes(ch, value_axis=not labels)
        _legend(ch, legend if legend is not None else len(series) > 1)
        if line:
            _add_line_overlay(gf, cats, line, cols[len(series) % len(cols)] if not colors else C["red"])
        return gf

    def line_chart(self, s, x, y, w, h, cats, series, colors=None, fmt='#,##0', end_labels=True, legend=None):
        """Line chart: 2.25 pt lines with round markers, a value label on the last point only.
        end_labels = True (all series), False, or a list of series names to label (skip ones whose ends would collide)."""
        cd = CategoryChartData(); cd.categories = cats
        for n, v in series.items(): cd.add_series(n, v)
        gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), cd); gf.name = "Chart"
        ch = gf.chart; ch.has_title = False; cols = colors or SERIES
        for i, ser in enumerate(ch.plots[0].series):
            c = rgb(cols[i % len(cols)])
            ser.format.line.color.rgb = c; ser.format.line.width = Pt(2.25); ser.smooth = False
            ser.marker.style = 8; ser.marker.size = 5
            ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = c; ser.marker.format.line.color.rgb = c
            if end_labels is True or (isinstance(end_labels, (list, tuple)) and ser.name in end_labels):
                last = ser.points[len(cats) - 1]; dl = last.data_label
                dl.has_text_frame = False; dl.position = XL_LABEL_POSITION.RIGHT
                dl.font.size = Pt(9); dl.font.name = FONT; dl.font.color.rgb = c; dl.font.bold = True
                _show_value(dl, fmt)
        _style_axes(ch, value_axis=True)
        _legend(ch, legend if legend is not None else len(series) > 1)
        return gf

    def doughnut_chart(self, s, x, y, w, h, cats, values, colors=None, fmt='#,##0', hole=55):
        """Doughnut for at most 4-5 shares (template: industries by category, outlay by component)."""
        cd = CategoryChartData(); cd.categories = cats; cd.add_series("Share", values)
        gf = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(x), Inches(y), Inches(w), Inches(h), cd); gf.name = "Chart"
        ch = gf.chart; ch.has_title = False; cols = colors or SERIES
        for i, pt in enumerate(ch.plots[0].series[0].points):
            pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(cols[i % len(cols)])
            pt.format.line.color.rgb = rgb(C["white"]); pt.format.line.width = Pt(1.5)
        hs = ch.plots[0]._element.find(qn("c:holeSize"))
        if hs is not None: hs.set("val", str(hole))
        p = ch.plots[0]; p.has_data_labels = True; dl = p.data_labels
        dl.number_format = fmt; dl.number_format_is_linked = False
        dl.font.size = Pt(9); dl.font.name = FONT; dl.font.color.rgb = rgb(C["white"]); dl.font.bold = True
        dl.show_value = True; dl.show_category_name = False; dl.show_percentage = False
        _legend(ch, True, XL_LEGEND_POSITION.RIGHT)
        return gf


# ---- chart and table helpers ---------------------------------------------------------------
def _font(el_txpr_owner, size=9, color=C["text"]):
    el_txpr_owner.font.size = Pt(size); el_txpr_owner.font.name = FONT; el_txpr_owner.font.color.rgb = rgb(color)


def _style_axes(ch, value_axis=False):
    ch.font.name = FONT; ch.font.size = Pt(9); ch.font.color.rgb = rgb(C["text"])
    try:
        ca = ch.category_axis
        ca.format.line.color.rgb = rgb(C["axis"]); ca.has_major_gridlines = False
        _font(ca.tick_labels, 9)
        tlp = ca._element.find(qn("c:tickLblPos"))
        if tlp is None:
            tlp = ca._element.makeelement(qn("c:tickLblPos"), {}); ca._element.find(qn("c:axPos")).addnext(tlp)
        tlp.set("val", "low")  # category labels stay below the plot even with negative bars
        va = ch.value_axis
        va.has_major_gridlines = False; va.has_minor_gridlines = False
        _font(va.tick_labels, 8)
        if value_axis:
            va.format.line.color.rgb = rgb(C["axis"])
        else:
            va.visible = False
    except (ValueError, AttributeError):
        pass


def _legend(ch, show, pos=XL_LEGEND_POSITION.BOTTOM):
    ch.has_legend = bool(show)
    if show:
        ch.legend.position = pos; ch.legend.include_in_layout = False
        _font(ch.legend, 9)


def _show_value(dl, fmt):
    d = dl._get_or_add_dLbl()
    nf = d.find(qn("c:numFmt"))
    if nf is None:
        nf = d.makeelement(qn("c:numFmt"), {}); d.insert(1, nf)
    nf.set("formatCode", fmt); nf.set("sourceLinked", "0")
    for tag in ("c:showLegendKey", "c:showVal", "c:showCatName", "c:showSerName", "c:showPercent", "c:showBubbleSize"):
        e = d.find(qn(tag))
        if e is None: e = d.makeelement(qn(tag), {}); d.append(e)
        e.set("val", "1" if tag == "c:showVal" else "0")


def _add_line_overlay(gf, cats, line, color):
    """Add a line series on the same axes as a bar chart (combo), e.g. demand over stacked supply."""
    from lxml import etree
    ch = gf.chart; plotArea = ch._chartSpace.chart.plotArea
    bar = plotArea.find(qn("c:barChart"))
    ax_ids = [a.get("val") for a in bar.findall(qn("c:axId"))]
    name, vals = next(iter(line.items()))
    n = len(bar.findall(qn("c:ser")))
    pts = "".join(f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in enumerate(vals))
    cpts = "".join(f'<c:pt idx="{i}"><c:v>{c}</c:v></c:pt>' for i, c in enumerate(cats))
    xml = (f'<c:lineChart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
           f'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><c:grouping val="standard"/><c:varyColors val="0"/>'
           f'<c:ser><c:idx val="{n}"/><c:order val="{n}"/><c:tx><c:v>{name}</c:v></c:tx>'
           f'<c:spPr><a:ln w="28575" cap="rnd"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill><a:round/></a:ln></c:spPr>'
           f'<c:marker><c:symbol val="circle"/><c:size val="5"/><c:spPr><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
           f'<a:ln><a:solidFill><a:srgbClr val="{color}"/></a:solidFill></a:ln></c:spPr></c:marker>'
           f'<c:cat><c:strLit><c:ptCount val="{len(cats)}"/>{cpts}</c:strLit></c:cat>'
           f'<c:val><c:numLit><c:formatCode>General</c:formatCode><c:ptCount val="{len(vals)}"/>{pts}</c:numLit></c:val>'
           f'<c:smooth val="0"/></c:ser><c:marker val="1"/>'
           f'<c:axId val="{ax_ids[0]}"/><c:axId val="{ax_ids[1]}"/></c:lineChart>')
    bar.addnext(etree.fromstring(xml))


def _cell_borders(cell, color=C["white"], w=31750):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None: tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat", "cmpd": "sng", "algn": "ctr"})
        sf = ln.makeelement(qn("a:solidFill"), {}); c = sf.makeelement(qn("a:srgbClr"), {"val": color}); sf.append(c); ln.append(sf)
        tcPr.append(ln)
    # schema order: borders must come before the cell fill
    fill = tcPr.find(qn("a:solidFill"))
    if fill is not None:
        tcPr.remove(fill); tcPr.append(fill)
